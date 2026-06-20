"""
Prototipo de reporte automático — sección "DEPARTAMENTO SANTIAGO".

Objetivo: demostrar con datos reales que el reporte de gastos que los contadores
arman a mano en Excel es una función DETERMINISTA de:
    (a) movimientos por código de cuenta con monto mensual  -> lo que /ledger-entries produce
    (b) un mapeo  código-de-cuenta -> bucket del reporte     -> el único artefacto manual

Estrategia de validación (cero credenciales, solo el propio Excel como ground truth):
  1. Leer las LÍNEAS DE DETALLE de la sección "Detalle de los gastos / Departamento Santiago"
     (cada línea = un código de cuenta con 12 montos mensuales). Esto hace de "ledger".
  2. Reconstruir los subtotales (Cuentas Básicas, Sueldos, Vehículos, Varios) y el
     total del centro de costo SUMANDO el detalle desde cero.
  3. Cuadrar lo reconstruido contra los totales IMPRESOS en el Excel y contra la
     línea-resumen del bloque EGRESOS (fila 78). Si todo da diff=0 -> feasibility probada.
"""
import sys
import openpyxl

XLSX = r"H:\My Drive\Emprendimientos\Analisis gastos Edu\Edu\Gastos EAG 12-2025.xlsx"
MONTHS = list(range(3, 15))  # columnas C..N = Ene..Dic
TOTAL_COL = 15               # columna O = TOTAL anual

# Estructura de la sección Depto Santiago en la hoja Resúmen (rango de detalle por subgrupo).
# (primera_fila_detalle, ultima_fila_detalle, fila_subtotal_impreso, nombre)
SUBGROUPS = [
    (130, 136, 137, "Cuentas Básicas"),
    (140, 142, 143, "Sueldos Brutos"),
    (146, 153, 154, "Vehículos"),
    (157, 164, 166, "Varios"),
]
COST_CENTER_TOTAL_ROW = 168       # "TOTAL DEPTO SANTIAGO"
SUMMARY_LINE_ROW = 78             # "DEPARTAMENTO SANTIAGO" en el bloque resumen EGRESOS


def num(v):
    return v if isinstance(v, (int, float)) else 0.0


def row_months(ws, r):
    return [num(ws.cell(r, c).value) for c in MONTHS]


def add(a, b):
    return [x + y for x, y in zip(a, b)]


def fmt(vec):
    return " ".join(f"{int(round(v)):>12,}" for v in vec)


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Resúmen"]

    print("=" * 110)
    print("PROTOTIPO REPORTE AUTOMÁTICO — Centro de costo: DEPARTAMENTO SANTIAGO  (año 2025)")
    print("Fuente: líneas de detalle del propio Excel como stand-in del ledger por cuenta.")
    print("=" * 110)

    cost_center = [0.0] * 12
    all_ok = True

    for first, last, sub_row, name in SUBGROUPS:
        computed = [0.0] * 12
        lines = 0
        for r in range(first, last + 1):
            label = ws.cell(r, 2).value
            if label is None or str(label).strip().upper().startswith("TOTAL"):
                continue
            vals = row_months(ws, r)
            if any(vals) or ws.cell(r, 1).value is not None:
                computed = add(computed, vals)
                lines += 1
        printed = row_months(ws, sub_row)
        diff = [round(c - p) for c, p in zip(computed, printed)]
        ok = not any(diff)
        all_ok &= ok
        cost_center = add(cost_center, computed)
        print(f"\n— {name}  ({lines} cuentas de detalle)  -> subtotal {'OK ✓' if ok else 'MISMATCH ✗'}")
        print(f"    reconstruido: {fmt(computed)}")
        if not ok:
            print(f"    impreso     : {fmt(printed)}")
            print(f"    diff        : {fmt(diff)}")

    # Total del centro de costo: reconstruido vs impreso (fila 168) vs línea-resumen (fila 78)
    printed_cc = row_months(ws, COST_CENTER_TOTAL_ROW)
    printed_summary = row_months(ws, SUMMARY_LINE_ROW)
    diff_cc = [round(c - p) for c, p in zip(cost_center, printed_cc)]
    diff_sum = [round(c - p) for c, p in zip(cost_center, printed_summary)]

    print("\n" + "-" * 110)
    print("TOTAL DEPARTAMENTO SANTIAGO")
    print(f"    reconstruido         : {fmt(cost_center)}   anual={int(round(sum(cost_center))):,}")
    print(f"    impreso (fila 168)   : {fmt(printed_cc)}")
    print(f"    impreso resumen (f78): {fmt(printed_summary)}")
    print(f"    diff vs fila 168     : {'OK ✓' if not any(diff_cc) else fmt(diff_cc)}")
    print(f"    diff vs resumen f78  : {'OK ✓' if not any(diff_sum) else fmt(diff_sum)}")

    all_ok &= not any(diff_cc) and not any(diff_sum)
    print("\n" + "=" * 110)
    print(f"RESULTADO: {'TODO CUADRA (diff=0) — feasibility probada ✓' if all_ok else 'HAY DESCUADRES — revisar ✗'}")
    print("=" * 110)
    return 0 if all_ok else 1


if __name__ == "__main__":
    sys.exit(main())
