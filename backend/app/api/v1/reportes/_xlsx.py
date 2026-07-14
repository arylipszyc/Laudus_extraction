"""Toolkit openpyxl reutilizable para reportes xlsx (extraído en Story 13.1).

Encapsula el cursor de fila + estilos que el reporte de EAG (`report_builder.py`)
tiene como closures internas. **El reporte de EAG NO se migra a este toolkit**: se
deja intacto para garantizar su output byte-idéntico (13.1 AC4 / NFR20). Queda
disponible para una refactor futura de EAG.

`ReportSheet` maneja un cursor de fila sobre una hoja con:
  columna A = "Detalle" (etiqueta) · N columnas de valor · (opcional) columna TOTAL.
Subtotales/totales se escriben como FÓRMULAS Excel para que recalculen si el
contador edita celdas.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LIGHT = PatternFill("solid", fgColor="D9D9D9")   # subtotales
DARK = PatternFill("solid", fgColor="A6A6A6")    # totales
WARN = PatternFill("solid", fgColor="FFF2CC")    # notas/limitaciones
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=12)
BIG = Font(bold=True, size=14)
ITALIC = Font(italic=True, size=9)
MONEY = "#,##0;-#,##0"                            # negativos sin rojo
BOTTOM = Border(bottom=Side(style="thin", color="808080"))


class ReportSheet:
    """Cursor de escritura sobre una worksheet con columnas de valor fijas."""

    def __init__(self, ws, n_value_cols: int, *, with_total: bool = True):
        self.ws = ws
        self.n = n_value_cols
        self.with_total = with_total
        # Letras de las columnas de valor (empiezan en B = col 2).
        self.LC = [get_column_letter(2 + j) for j in range(n_value_cols)]
        self.total_col = 2 + n_value_cols  # solo se usa si with_total
        self.r = 1
        ws.sheet_view.showGridLines = False

    def title(self, text: str, *, big: bool = False):
        self.ws.cell(self.r, 1, text).font = BIG if big else TITLE
        self.r += 1

    def note(self, text: str, *, italic: bool = True, fill=None, bold: bool = False):
        cell = self.ws.cell(self.r, 1, text)
        cell.font = ITALIC if italic else (BOLD if bold else Font())
        if fill:
            cell.fill = fill
        self.r += 1

    def blank(self):
        self.r += 1

    def header_row(self, headers: list[str], *, label_title: str = "Detalle"):
        cols = [(1, label_title)] + [(2 + j, h) for j, h in enumerate(headers)]
        if self.with_total:
            cols.append((self.total_col, "TOTAL"))
        for c, txt in cols:
            cell = self.ws.cell(self.r, c, txt)
            cell.font, cell.border = BOLD, BOTTOM
            if c > 1:
                cell.alignment = Alignment(horizontal="center")
        self.r += 1

    def _row_total_formula(self, row: int):
        return f"=SUM({self.LC[0]}{row}:{self.LC[-1]}{row})" if self.n else None

    def row(self, label, *, values=None, formulas=None, bold=False, indent=False,
            fill=None, total=None):
        """Escribe una fila. `values` = números (se redondean); `formulas` = strings
        Excel. `total`: None→auto =SUM (si with_total), string→fórmula, False→sin total."""
        cur = self.r
        cell = self.ws.cell(cur, 1, ("   " if indent else "") + str(label))
        if bold:
            cell.font = BOLD
        for j in range(self.n):
            c = self.ws.cell(cur, 2 + j)
            if formulas is not None:
                c.value = formulas[j]
            elif values is not None:
                c.value = round(values[j])
            c.number_format = MONEY
            if bold:
                c.font = BOLD
        if self.with_total and total is not False:
            tc = self.ws.cell(cur, self.total_col)
            tc.value = total if isinstance(total, str) else self._row_total_formula(cur)
            tc.number_format = MONEY
            if bold:
                tc.font = BOLD
        if fill:
            last = self.total_col if self.with_total else 1 + self.n
            for c in range(1, last + 1):
                self.ws.cell(cur, c).fill = fill
        self.r += 1
        return cur

    def sum_formula(self, first: int, last: int):
        """Fórmula por columna =SUM(col{first}:col{last}) para un subtotal de rango."""
        return [f"=SUM({self.LC[j]}{first}:{self.LC[j]}{last})" for j in range(self.n)]

    def add_formula(self, rows: list[int]):
        """Fórmula por columna =c1+c2+... para un total de filas no contiguas."""
        return ["=" + "+".join(f"{self.LC[j]}{rr}" for rr in rows) for j in range(self.n)]

    def set_widths(self, label_width: int, value_width: int):
        self.ws.column_dimensions["A"].width = label_width
        n = self.n + (1 if self.with_total else 0)
        for j in range(n):
            self.ws.column_dimensions[get_column_letter(2 + j)].width = value_width
