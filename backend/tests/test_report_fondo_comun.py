"""Tests del reporte del Fondo Común (RUT2 · FFCC/JAB) — Story 13.1.

Cubre:
  · `distribution_rows_via_beancount`: neteo de asientos wash, split familiares/
    operativas, invariante inicial+retiros+repartos==cierre, saldos fieles a Laudus.
  · `build_report_rut2`: dos hojas (Gastos + Distribuciones), agrupación mecánica
    por encabezado, TC lumpeada marcada en el cuerpo (AC2), marca de limitación
    honesta obligatoria (AC3b), slots as-if inversiones/offshore.
  · Endpoint GET /reportes/fondo-comun: 200 xlsx, 422 start>end, RBAC.
"""
import io
from datetime import date

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.api.v1.reportes.report_builder_rut2 import build_report_rut2
from backend.app.auth.service import create_jwt
from backend.app.dependencies import get_ledger_service
from backend.app.middleware import add_middleware
from backend.app.services.bql_queries import (
    distribution_rows_via_beancount,
    report_rows_via_beancount,
)
from backend.app.services.ledger_service import LedgerService

# Ledger FFCC/JAB autocontenido: gastos por encabezado (41/43 · 81/87 con TC),
# CxC familiares (AAG a favor, SAG debe) + operativa (FondoFijo), y un asiento WASH
# (+X/−X a la misma cuenta, neto 0) que NO debe inflar retiros/repartos.
FONDO_LEDGER = """\
option "operating_currency" "CLP"
1900-01-01 commodity CLP

2020-01-01 open Assets:FFCC:Banco-111005F CLP
  code: "111005F"
2020-01-01 open Equity:FFCC:Utilidades CLP

2020-01-01 open Expenses:FFCC:GastosGenerales-410001 CLP
  code: "410001"
  laudus_account_name: "Gastos Generales FFCC"
  laudus_categoria2: "GASTOS GENERALES"
2020-01-01 open Expenses:FFCC:Oficina-430001 CLP
  code: "430001"
  laudus_account_name: "Oficina FFCC"
  laudus_categoria2: "GASTOS FIJOS OFICINA"
2020-01-01 open Expenses:JAB:CasasMantencion-810001 CLP
  code: "810001"
  laudus_account_name: "Mantencion Casa"
  laudus_categoria2: "MANTENCION CASAS JAB"
2020-01-01 open Expenses:JAB:TarjetaJab-871005 CLP
  code: "871005"
  laudus_account_name: "JAB Mastercard Visa Amex"
  laudus_categoria2: "GASTOS PERSONALES"
  laudus_categoria3: "JAB - GASTOS PERSONALES"

2020-01-01 open Assets:FFCC:RetirosAag-115021 CLP
  code: "115021"
  laudus_account_name: "Retiros AAG"
  laudus_categoria2: "ACTIVOS CORRIENTES"
  laudus_categoria3: "CUENTAS POR COBRAR"
2020-01-01 open Assets:FFCC:RetirosSag-115025 CLP
  code: "115025"
  laudus_account_name: "Retiros SAG"
  laudus_categoria2: "ACTIVOS CORRIENTES"
  laudus_categoria3: "CUENTAS POR COBRAR"
2020-01-01 open Assets:FFCC:FondoFijo-115003 CLP
  code: "115003"
  laudus_account_name: "Fondo Fijo"
  laudus_categoria2: "ACTIVOS CORRIENTES"
  laudus_categoria3: "CUENTAS POR COBRAR"

; ── Gastos del período 2025 ──
2025-03-10 * "Gasto general FFCC"
  Expenses:FFCC:GastosGenerales-410001    40000 CLP
  Assets:FFCC:Banco-111005F              -40000 CLP
2025-03-11 * "Oficina FFCC"
  Expenses:FFCC:Oficina-430001    10000 CLP
  Assets:FFCC:Banco-111005F      -10000 CLP
2025-04-05 * "Mantencion casa JAB"
  Expenses:JAB:CasasMantencion-810001    80000 CLP
  Assets:FFCC:Banco-111005F             -80000 CLP
2025-04-06 * "Pago TC JAB (lump)"
  Expenses:JAB:TarjetaJab-871005    55000 CLP
  Assets:FFCC:Banco-111005F        -55000 CLP

; ── Distribuciones ──
; AAG: inicial 2024 +1.000.000; período retiro +300.000, reparto -2.000.000 → cierre -700.000 (a favor)
2024-06-01 * "Retiro AAG 2024"
  Assets:FFCC:RetirosAag-115021    1000000 CLP
  Assets:FFCC:Banco-111005F       -1000000 CLP
2025-02-01 * "Retiro AAG"
  Assets:FFCC:RetirosAag-115021     300000 CLP
  Assets:FFCC:Banco-111005F        -300000 CLP
2025-12-31 * "Reparto utilidad AAG (asignada, no retirada)"
  Assets:FFCC:RetirosAag-115021   -2000000 CLP
  Equity:FFCC:Utilidades           2000000 CLP

; SAG: inicial 2024 +500.000 → cierre +500.000 (debe). WASH 2025 no debe inflar flujos.
2024-05-01 * "Retiro SAG 2024"
  Assets:FFCC:RetirosSag-115025     500000 CLP
  Assets:FFCC:Banco-111005F        -500000 CLP
2025-06-01 * "Wash reclasificacion SAG (neto 0)"
  Assets:FFCC:RetirosSag-115025     900000 CLP
  Assets:FFCC:RetirosSag-115025    -900000 CLP

; Operativa
2025-03-01 * "Fondo fijo"
  Assets:FFCC:FondoFijo-115003     5000 CLP
  Assets:FFCC:Banco-111005F       -5000 CLP
"""

START, END = date(2025, 1, 1), date(2025, 12, 31)


def _ledger(tmp_path):
    main = tmp_path / "main.beancount"
    main.write_text(FONDO_LEDGER, encoding="utf-8")
    return LedgerService(str(main))


def _dist(tmp_path):
    return distribution_rows_via_beancount(
        _ledger(tmp_path), START.isoformat(), END.isoformat(), group="FondoComun")


def _by_code(rows):
    return {r["code"]: r for r in rows}


# ── distribution_rows_via_beancount ──────────────────────────────────────────


def test_dist_familiar_vs_operational_split(tmp_path):
    rows = _dist(tmp_path)
    by = _by_code(rows)
    assert by["115021"]["operational"] is False   # RetirosAag = familiar
    assert by["115025"]["operational"] is False   # RetirosSag = familiar
    assert by["115003"]["operational"] is True    # FondoFijo = operativa
    # Familiares primero (orden del servicio), operativas al final.
    ops = [r["operational"] for r in rows]
    assert ops == sorted(ops)


def test_dist_wash_netting(tmp_path):
    """El asiento wash (+900k/−900k a RetirosSag en una tx) NO debe inflar
    retiros ni repartos: netea a 0 dentro del asiento."""
    sag = _by_code(_dist(tmp_path))["115025"]
    assert sag["retiros"] == 0.0
    assert sag["repartos"] == 0.0
    assert sag["saldo_cierre"] == 500000.0   # solo el retiro 2024 (inicial)


def test_dist_invariant_and_period_flows(tmp_path):
    """Invariante inicial + retiros + repartos == cierre; los flujos son del período."""
    aag = _by_code(_dist(tmp_path))["115021"]
    assert aag["saldo_inicial"] == 1000000.0
    assert aag["retiros"] == 300000.0
    assert aag["repartos"] == -2000000.0
    assert aag["saldo_cierre"] == -700000.0
    assert aag["saldo_inicial"] + aag["retiros"] + aag["repartos"] == aag["saldo_cierre"]


def test_dist_matches_laudus_balance(tmp_path):
    """Los saldos son fieles al ledger (== balance_sheet del grupo a la fecha)."""
    from backend.app.services.bql_queries import balance_sheet_via_beancount
    bs = balance_sheet_via_beancount(_ledger(tmp_path), "FondoComun", date_to=END.isoformat())
    bs_by = {r["account_number"]: r for r in bs["data"]}
    for d in _dist(tmp_path):
        row = bs_by[d["code"]]
        signed = row["debit_balance"] - row["credit_balance"]
        assert round(signed) == round(d["saldo_cierre"]), d["code"]


# ── build_report_rut2 ────────────────────────────────────────────────────────


def _build(tmp_path):
    ledger = _ledger(tmp_path)
    rows = report_rows_via_beancount(ledger, START.isoformat(), END.isoformat(), group="FondoComun")
    dist = distribution_rows_via_beancount(ledger, START.isoformat(), END.isoformat(), group="FondoComun")
    return build_report_rut2(START, END, rows, dist)


def test_build_two_sheets_and_valid_xlsx(tmp_path):
    data = _build(tmp_path)
    assert isinstance(data, bytes) and data[:2] == b"PK"
    wb = load_workbook(io.BytesIO(data))
    assert "Gastos FFCC-JAB" in wb.sheetnames
    assert "Distribuciones" in wb.sheetnames


def _text(ws):
    return " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def test_gastos_headers_and_tc_marked(tmp_path):
    ws = load_workbook(io.BytesIO(_build(tmp_path)))["Gastos FFCC-JAB"]
    t = _text(ws)
    for enc in ("GASTOS GENERALES", "GASTOS FIJOS OFICINA", "MANTENCION CASAS JAB", "GASTOS PERSONALES"):
        assert enc in t, enc
    assert "TOTAL GASTOS FONDO COMÚN" in t
    # AC2: TC 871005 lumpeada marcada en el CUERPO (no nota al pie).
    assert "871005" in t and "TC lump" in t


def test_distribuciones_limitation_and_slots(tmp_path):
    ws = load_workbook(io.BytesIO(_build(tmp_path)))["Distribuciones"]
    t = _text(ws)
    # AC3b: marca de limitación honesta en el cuerpo.
    assert "no es determinable" in t.lower()
    assert "FAMILIARES" in t and "OPERATIVAS" in t
    # Etiquetas de signo (Valentina).
    assert "A favor" in t and "Debe al fondo" in t
    # Slots as-if inversiones/offshore (Ary 2026-07-13).
    assert "PENDIENTE DE CARGA" in t and "offshore" in t.lower()


# ── Endpoint GET /reportes/fondo-comun ───────────────────────────────────────


def _app(tmp_path):
    from backend.app.api.v1.reportes.router import router as reportes_router
    svc = _ledger(tmp_path)
    app = FastAPI()
    add_middleware(app)
    app.include_router(reportes_router, prefix="/api/v1")
    app.dependency_overrides[get_ledger_service] = lambda: svc
    return TestClient(app, raise_server_exceptions=False)


def test_endpoint_ok_xlsx(tmp_path):
    client = _app(tmp_path)
    resp = client.get("/api/v1/reportes/fondo-comun",
                      params={"start": "2025-01-01", "end": "2025-12-31"},
                      cookies={"access_token": create_jwt(email="c@t.com", role="contador")})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert resp.content[:2] == b"PK"
    assert "reporte_fondo_comun" in resp.headers.get("content-disposition", "")


def test_endpoint_start_after_end_422(tmp_path):
    client = _app(tmp_path)
    resp = client.get("/api/v1/reportes/fondo-comun",
                      params={"start": "2025-12-31", "end": "2025-01-01"},
                      cookies={"access_token": create_jwt(email="c@t.com", role="contador")})
    assert resp.status_code == 422


def test_endpoint_rbac_family_denied(tmp_path):
    client = _app(tmp_path)
    resp = client.get("/api/v1/reportes/fondo-comun",
                      params={"start": "2025-01-01", "end": "2025-12-31"},
                      cookies={"access_token": create_jwt(email="f@t.com", role="family")})
    assert resp.status_code == 403
