"""Tests for report_rows_via_beancount — migración del reporte a Beancount (#4 / 9.11).

Valida que las filas estilo `ledger_final` derivadas del ledger Beancount tengan la
forma que `report_builder` consume, cubran TODO el grupo EAG (matriz + hijas, Story
11.1 — antes "todas las entidades"), hagan el split debit/credit por signo, y respeten
el rango de fechas. Cierra con una prueba de integración: las filas alimentan
`build_report` y producen un xlsx válido.
"""
from datetime import date

from openpyxl import load_workbook
import io

from backend.app.api.v1.reportes.report_builder import build_report
from backend.app.services.bql_queries import report_rows_via_beancount
from backend.app.services.ledger_service import LedgerService


MINI_LEDGER = """\
option "operating_currency" "CLP"
1900-01-01 commodity CLP

2020-01-01 open Assets:EAG:Bancos:TestBank-111005 CLP
  code: "111005"
  laudus_account_name: "Banco Test EAG"
  laudus_categoria1: "ACTIVO EAG"
  laudus_categoria2: "ACTIVOS CORRIENTES"
  laudus_categoria3: "DISPONIBLE - EAG"
2020-01-01 open Income:EAG:Ventas-411005 CLP
  code: "411005"
2020-01-01 open Expenses:EAG:Gastos-511005 CLP
  code: "511005"
2020-01-01 open Assets:Jocelyn:Bancos:TestBankJ-610005 CLP
  code: "610005"
  laudus_account_name: "Banco Test Jocelyn"

2024-03-15 * "Pago cliente"
  Assets:EAG:Bancos:TestBank-111005   100000 CLP
  Income:EAG:Ventas-411005           -100000 CLP

2024-06-20 * "Compra insumos"
  Expenses:EAG:Gastos-511005           30000 CLP
  Assets:EAG:Bancos:TestBank-111005   -30000 CLP

2024-04-10 * "Gasto Jocelyn"
  Assets:Jocelyn:Bancos:TestBankJ-610005   50000 CLP
  Income:EAG:Ventas-411005                -50000 CLP
"""

# Las claves mínimas que report_builder lee de cada fila.
ROW_KEYS = {"date", "accountnumber", "accountName",
            "Categoria1", "Categoria2", "Categoria3", "debit", "credit"}


def _ledger(tmp_path):
    main = tmp_path / "main.beancount"
    main.write_text(MINI_LEDGER, encoding="utf-8")
    return LedgerService(str(main))


def test_rows_carry_report_builder_keys(tmp_path):
    rows = report_rows_via_beancount(_ledger(tmp_path))
    assert rows
    for r in rows:
        assert ROW_KEYS <= set(r.keys())


def test_includes_whole_eag_group(tmp_path):
    """Story 11.1: el reporte cubre TODO el grupo EAG (matriz + hijas) — a
    diferencia de ledger_entries_via_beancount no es per-entity, pero tampoco
    "todas las entidades": lo que no pertenece al grupo EAG queda fuera."""
    rows = report_rows_via_beancount(_ledger(tmp_path))
    numbers = {r["accountnumber"] for r in rows}
    assert "111005" in numbers   # EAG
    assert "610005" in numbers   # Jocelyn


def test_excludes_foreign_namespace_keeps_legacy_equity(tmp_path):
    """Story 11.1 AC2: una cuenta FFCC (fuera del grupo EAG) no aparece en las
    filas del reporte; los namespaces Equity sin entidad (legacy EAG) siguen."""
    extra = """\

2020-01-01 open Assets:FFCC:Test-410001 CLP
  code: "410001"
2020-01-01 open Equity:FFCC:Apertura CLP
2020-01-01 open Equity:Apertura:TarjetasSinDetalle CLP

2024-05-05 * "Aporte FFCC"
  Assets:FFCC:Test-410001    200000 CLP
  Equity:FFCC:Apertura      -200000 CLP

2024-05-07 * "Apertura TC sin detalle"
  Assets:EAG:Bancos:TestBank-111005       -80000 CLP
  Equity:Apertura:TarjetasSinDetalle       80000 CLP
"""
    main = tmp_path / "main.beancount"
    main.write_text(MINI_LEDGER + extra, encoding="utf-8")
    rows = report_rows_via_beancount(LedgerService(str(main)))
    numbers = {r["accountnumber"] for r in rows}
    assert "410001" not in numbers                       # FFCC excluida (AC2)
    # Guard TRAP #1: la pata Equity:Apertura sí entra (2 filas del 2024-05-07).
    legacy_day = [r for r in rows if r["date"] == "2024-05-07"]
    assert len(legacy_day) == 2


def test_debit_credit_split_by_sign(tmp_path):
    """Posting positivo → debit; negativo → credit. Es el signo que report_builder
    convierte en monto ((de-cr) gasto/activo, (cr-de) ingreso)."""
    rows = report_rows_via_beancount(_ledger(tmp_path))
    # El cargo al banco (+100000) y el ingreso contrapartida (-100000) del 2024-03-15.
    bank_in = next(r for r in rows if r["accountnumber"] == "111005" and r["date"] == "2024-03-15")
    assert bank_in["debit"] == 100000.0 and bank_in["credit"] == 0.0
    income = next(r for r in rows if r["accountnumber"] == "411005" and r["date"] == "2024-03-15")
    assert income["credit"] == 100000.0 and income["debit"] == 0.0


def test_date_range_bounds_rows(tmp_path):
    """date_from/date_to acotan los postings devueltos."""
    rows = report_rows_via_beancount(_ledger(tmp_path), date_from="2024-05-01")
    dates = {r["date"] for r in rows}
    assert dates == {"2024-06-20"}            # solo la compra de insumos
    rows2 = report_rows_via_beancount(_ledger(tmp_path), date_from="2024-04-01", date_to="2024-04-30")
    assert {r["date"] for r in rows2} == {"2024-04-10"}


def test_metadata_join_and_fallbacks(tmp_path):
    """accountnumber=code, categorías desde el Open; sin meta → accountName cae al
    nombre de cuenta y categorías quedan ''."""
    rows = report_rows_via_beancount(_ledger(tmp_path))
    bank = next(r for r in rows if r["accountnumber"] == "111005")
    assert bank["accountName"] == "Banco Test EAG"
    assert bank["Categoria1"] == "ACTIVO EAG"
    income = next(r for r in rows if r["accountnumber"] == "411005")
    assert income["Categoria1"] == ""        # el Open de Income no trae categoría
    assert income["accountName"]             # cae al nombre de cuenta, no vacío


def test_rows_feed_build_report(tmp_path):
    """Integración: las filas derivadas de Beancount alimentan build_report y
    producen un xlsx válido (mismo contrato que el path Sheets)."""
    rows = report_rows_via_beancount(_ledger(tmp_path))
    data = build_report(date(2024, 1, 1), date(2024, 12, 31), lambda _name: rows)
    assert isinstance(data, bytes) and data[:2] == b"PK"   # firma ZIP/xlsx
    wb = load_workbook(io.BytesIO(data))
    assert wb.active.title == "Reporte"


def test_report_itemizes_daughter_accounts():
    """El reporte desglosa cuenta por cuenta dentro de cada hija (sección DETALLE HIJAS),
    como el workbook del contador — no solo el subtotal."""
    daughter = "EGRESOS JOCELYN AVAYU DEUTSCH"
    rows = [
        {"date": "2026-03-10", "accountnumber": "690003", "accountName": "Jocelyn - PAC Seguros",
         "Categoria1": daughter, "Categoria2": "", "Categoria3": "", "debit": 1177, "credit": 0},
        {"date": "2026-03-15", "accountnumber": "690001", "accountName": "Jocelyn - Remesa",
         "Categoria1": daughter, "Categoria2": "", "Categoria3": "", "debit": 2500000, "credit": 0},
    ]
    data = build_report(date(2026, 1, 1), date(2026, 6, 30), lambda _name: rows)
    ws = load_workbook(io.BytesIO(data)).active
    labels = [str(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value]
    text = " ".join(labels)
    assert any("DETALLE DE GASTOS" in l and "HIJAS" in l for l in labels)
    assert "690003" in text and "690001" in text          # ambas cuentas itemizadas
    assert "Jocelyn - PAC Seguros" in text                 # con su nombre
