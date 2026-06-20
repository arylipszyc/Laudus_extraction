"""Tests de cuentas pendientes — Story 10.3 (promoción desde cuarentena).

Cubre el contrato de los AC contra un ledger fixture en `tmp_path` (mismo patrón que
test_dashboard_beancount): listar pendientes (AC1), sugerencia por prefijo (AC2),
promoción con bean-check verde (AC3), rollback con bean-check rojo (AC3), reconciliación
de cuarentena tras promover (AC4), el reporte deja de marcar la cuenta como sin categorizar
(AC6), y RBAC (AC8).
"""
import io
from datetime import date

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.api.v1.cuentas_pendientes.router import router as cp_router
from backend.app.auth.service import create_jwt
from backend.app.dependencies import get_ledger_service
from backend.app.middleware import add_middleware
from backend.app.services.beancount_promote import build_open_block
from backend.app.services.ledger_service import LedgerService

# ── Ledger fixture ────────────────────────────────────────────────────────────

ACCOUNTS = """\
2020-12-31 open Assets:EAG:Caja-111001 CLP
  code: "111001"
  laudus_account_name: "Caja $"
  laudus_categoria1: "ACTIVO EAG"
  laudus_categoria2: "ACTIVOS CORRIENTES"
  laudus_categoria3: "DISPONIBLE - EAG"

2020-12-31 open Expenses:EAG:Luz-413005 CLP
  code: "413005"
  laudus_account_name: "Luz Casa Sur"
  laudus_categoria1: "EGRESOS CASA SUR"
  laudus_categoria2: "Casa Sur"
  laudus_categoria3: "Servicios Basicos"

2020-12-31 open Expenses:EAG:Agua-413013 CLP
  code: "413013"
  laudus_account_name: "Agua Casa Sur"
  laudus_categoria1: "EGRESOS CASA SUR"
  laudus_categoria2: "Casa Sur"
  laudus_categoria3: "Servicios Basicos"
"""

PENDING = """\
2020-01-01 open Assets:EAG:PendingReview:Cuenta-413077 CLP
  code: "413077"
  pending_review: "TRUE"
"""

TXNS = """\
2025-03-10 * "Gasto nuevo casa sur" #pending-account
  id: "999"
  source: "laudus-erp"
  Assets:EAG:PendingReview:Cuenta-413077   150000 CLP
  Assets:EAG:Caja-111001                  -150000 CLP
"""

MAIN = """\
option "operating_currency" "CLP"
1900-01-01 commodity CLP
include "accounts.beancount"
include "pending.beancount"
include "txns.beancount"
"""


def _make_ledger(tmp_path, txns=TXNS):
    (tmp_path / "accounts.beancount").write_text(ACCOUNTS, encoding="utf-8")
    (tmp_path / "pending.beancount").write_text(PENDING, encoding="utf-8")
    (tmp_path / "txns.beancount").write_text(txns, encoding="utf-8")
    main = tmp_path / "main.beancount"
    main.write_text(MAIN, encoding="utf-8")
    return LedgerService(str(main))


def _make_app(svc) -> TestClient:
    app = FastAPI()
    add_middleware(app)
    app.include_router(cp_router, prefix="/api/v1")
    app.dependency_overrides[get_ledger_service] = lambda: svc
    return TestClient(app, raise_server_exceptions=False)


def _contador():
    return create_jwt(email="contador@test.com", role="contador")


def _family():
    return create_jwt(email="family@test.com", role="family")


def _admin():
    return create_jwt(email="admin@test.com", role="admin")


def _on(monkeypatch):
    monkeypatch.setenv("USE_BEANCOUNT_ENGINE_LEDGER", "true")
    monkeypatch.delenv("IMPORTER_GIT_ENABLED", raising=False)  # git no-op en tests
    monkeypatch.delenv("LEDGER_DIR", raising=False)            # refresh clone no-op


# ── AC1: listar pendientes ──────────────────────────────────────────────────


def test_lista_pendientes_con_monto(tmp_path, monkeypatch):
    _on(monkeypatch)
    client = _make_app(_make_ledger(tmp_path))
    resp = client.get("/api/v1/cuentas-pendientes/", cookies={"access_token": _contador()})
    assert resp.status_code == 200
    body = resp.json()
    assert len(body) == 1
    pa = body[0]
    assert pa["code"] == "413077"
    assert pa["pending_account"] == "Assets:EAG:PendingReview:Cuenta-413077"
    assert pa["monto_acumulado"] == 150000.0


# ── AC2: sugerencia por prefijo ─────────────────────────────────────────────


def test_sugerencia_cat1_cat2_por_prefijo_sin_cat3(tmp_path, monkeypatch):
    _on(monkeypatch)
    client = _make_app(_make_ledger(tmp_path))
    body = client.get("/api/v1/cuentas-pendientes/", cookies={"access_token": _contador()}).json()
    sug = body[0]["suggestion"]
    # 413xxx → hereda Cat1/Cat2 de las hermanas (413005/413013), Cat3 nunca se sugiere.
    assert sug["categoria1"] == "EGRESOS CASA SUR"
    assert sug["categoria2"] == "Casa Sur"
    assert sug["account"].startswith("Expenses:EAG:")
    assert sug["account"].endswith("-413077")


def test_sugerencia_vacia_sin_hermanas(tmp_path, monkeypatch):
    _on(monkeypatch)
    # Pending de un prefijo (999) que no tiene hermanas en el plan.
    pending = PENDING.replace("413077", "999001")
    txns = TXNS.replace("413077", "999001")
    (tmp_path / "accounts.beancount").write_text(ACCOUNTS, encoding="utf-8")
    (tmp_path / "pending.beancount").write_text(pending, encoding="utf-8")
    (tmp_path / "txns.beancount").write_text(txns, encoding="utf-8")
    main = tmp_path / "main.beancount"
    main.write_text(MAIN, encoding="utf-8")
    client = _make_app(LedgerService(str(main)))
    body = client.get("/api/v1/cuentas-pendientes/", cookies={"access_token": _contador()}).json()
    sug = body[0]["suggestion"]
    assert sug["categoria1"] == "" and sug["categoria2"] == "" and sug["account"] == ""


# ── AC3: promover (bean-check verde) ────────────────────────────────────────


def test_promover_escribe_open_y_refresca(tmp_path, monkeypatch):
    _on(monkeypatch)
    svc = _make_ledger(tmp_path)
    client = _make_app(svc)
    resp = client.post(
        "/api/v1/cuentas-pendientes/413077/promover",
        json={
            "categoria1": "EGRESOS CASA SUR",
            "categoria2": "Casa Sur",
            "categoria3": "Mantencion General",
            "laudus_account_name": "Gasto Nuevo Casa Sur",
        },
        cookies={"access_token": _contador()},
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["account"] == "Expenses:EAG:GastoNuevoCasaSur-413077"
    assert body["backfill_recommended"] is True
    # La cuenta quedó escrita en accounts.beancount con su metadata.
    written = (tmp_path / "accounts.beancount").read_text(encoding="utf-8")
    assert "Expenses:EAG:GastoNuevoCasaSur-413077" in written
    assert 'laudus_categoria3: "Mantencion General"' in written
    # Y el ledger en proceso ya la ve (AC6 — load() tras escribir).
    from beancount.core.data import Open
    accounts = {e.account for e in svc.entries() if isinstance(e, Open)}
    assert "Expenses:EAG:GastoNuevoCasaSur-413077" in accounts


def test_promover_cat3_obligatorio(tmp_path, monkeypatch):
    _on(monkeypatch)
    client = _make_app(_make_ledger(tmp_path))
    resp = client.post(
        "/api/v1/cuentas-pendientes/413077/promover",
        json={"categoria1": "EGRESOS CASA SUR", "categoria2": "Casa Sur", "categoria3": "  "},
        cookies={"access_token": _contador()},
    )
    assert resp.status_code == 422


def test_promover_code_no_pendiente_404(tmp_path, monkeypatch):
    _on(monkeypatch)
    client = _make_app(_make_ledger(tmp_path))
    resp = client.post(
        "/api/v1/cuentas-pendientes/111001/promover",
        json={"categoria1": "X", "categoria2": "Y", "categoria3": "Z"},
        cookies={"access_token": _contador()},
    )
    assert resp.status_code == 404


# ── AC3: rollback (bean-check rojo) ─────────────────────────────────────────


def test_promover_bean_check_rojo_hace_rollback(tmp_path, monkeypatch):
    _on(monkeypatch)
    svc = _make_ledger(tmp_path)
    client = _make_app(svc)
    original = (tmp_path / "accounts.beancount").read_text(encoding="utf-8")
    # Forzar bean-check rojo: override a una cuenta ya abierta → "Duplicate open".
    resp = client.post(
        "/api/v1/cuentas-pendientes/413077/promover",
        json={
            "categoria1": "EGRESOS CASA SUR", "categoria2": "Casa Sur", "categoria3": "X",
            "account": "Expenses:EAG:Luz-413005",
        },
        cookies={"access_token": _contador()},
    )
    assert resp.status_code == 422
    # No se commiteó: accounts.beancount quedó EXACTAMENTE como estaba (rollback).
    assert (tmp_path / "accounts.beancount").read_text(encoding="utf-8") == original


# ── AC4: reconciliación de cuarentena tras promover ─────────────────────────


def test_re_run_writer_no_re_emite_code_promovido(tmp_path):
    """Tras escribir el open (lo que hace promover), un re-run del writer saca el code
    de la cuarentena y postea las JEs a la cuenta real."""
    from pipeline.writers.beancount_writer import write_jes, _existing_pending_codes

    root = tmp_path
    target_dir = root / "imports" / "laudus"
    target_dir.mkdir(parents=True)
    accounts_path = root / "accounts.beancount"
    pending_path = root / "imports" / "_new-accounts-pending.beancount"
    accounts_path.write_text(ACCOUNTS, encoding="utf-8")

    rows = [
        {"journalentryid": "999", "journalentrynumber": "1", "date": "2025-03-10",
         "description": "Gasto", "accountnumber": "413077", "debit": 150000, "credit": 0},
        {"journalentryid": "999", "journalentrynumber": "1", "date": "2025-03-10",
         "description": "Gasto", "accountnumber": "111001", "debit": 0, "credit": 150000},
    ]

    # 1) Antes de promover: 413077 ausente del plan → cae a cuarentena.
    write_jes(rows, target_dir, accounts_path, pending_path, replace=True)
    assert "413077" in _existing_pending_codes(pending_path)
    month = (target_dir / "2025-03.beancount").read_text(encoding="utf-8")
    assert "PendingReview:Cuenta-413077" in month

    # 2) Promover = appendear el open al plan (mismo bloque que escribe el helper).
    accounts_path.write_text(
        ACCOUNTS + build_open_block(
            "Expenses:EAG:GastoNuevoCasaSur-413077", "413077",
            "Gasto Nuevo Casa Sur", "EGRESOS CASA SUR", "Casa Sur", "Mantencion",
        ),
        encoding="utf-8",
    )

    # 3) Re-run del writer: 413077 ya está en el índice → cuarentena se limpia + JE a cuenta real.
    write_jes(rows, target_dir, accounts_path, pending_path, replace=True)
    assert "413077" not in _existing_pending_codes(pending_path)
    month = (target_dir / "2025-03.beancount").read_text(encoding="utf-8")
    assert "Expenses:EAG:GastoNuevoCasaSur-413077" in month
    assert "PendingReview:Cuenta-413077" not in month


# ── AC6: el reporte deja de marcar la cuenta como "sin categorizar" ─────────


def test_pre_backfill_aparece_sin_categorizar(tmp_path, monkeypatch):
    _on(monkeypatch)
    # JE posteada a la cuenta de cuarentena (estado pre-backfill) → sin categoría.
    svc = _make_ledger(tmp_path)
    from backend.app.services.bql_queries import report_rows_via_beancount
    rows = report_rows_via_beancount(svc, "2025-01-01", "2025-12-31")
    by_code = {r["accountnumber"]: r for r in rows}
    assert by_code["413077"]["Categoria2"] == ""  # cuarentena no tiene categoría


def test_post_promover_y_backfill_queda_categorizada(tmp_path, monkeypatch):
    _on(monkeypatch)
    # Estado post-promoción+backfill: el open de 413077 existe con categorías y la JE
    # ya está posteada a la cuenta real (lo que deja el backfill).
    promoted_accounts = ACCOUNTS + build_open_block(
        "Expenses:EAG:GastoNuevoCasaSur-413077", "413077",
        "Gasto Nuevo Casa Sur", "EGRESOS CASA SUR", "Casa Sur", "Mantencion General",
    )
    backfilled_txn = """\
2025-03-10 * "Gasto nuevo casa sur"
  id: "999"
  source: "laudus-erp"
  Expenses:EAG:GastoNuevoCasaSur-413077   150000 CLP
  Assets:EAG:Caja-111001                 -150000 CLP
"""
    (tmp_path / "accounts.beancount").write_text(promoted_accounts, encoding="utf-8")
    (tmp_path / "pending.beancount").write_text("", encoding="utf-8")
    (tmp_path / "txns.beancount").write_text(backfilled_txn, encoding="utf-8")
    main = tmp_path / "main.beancount"
    main.write_text(MAIN, encoding="utf-8")
    svc = LedgerService(str(main))

    from backend.app.services.bql_queries import report_rows_via_beancount
    rows = report_rows_via_beancount(svc, "2025-01-01", "2025-12-31")
    by_code = {r["accountnumber"]: r for r in rows}
    # La plata cae en su Cat2, ya no en el balde sin categoría.
    assert by_code["413077"]["Categoria2"] == "Casa Sur"

    from backend.app.api.v1.reportes.report_builder import build_report
    out = build_report(date(2025, 1, 1), date(2025, 12, 31), lambda _n: rows)
    ws = load_workbook(io.BytesIO(out)).active
    labels = [str(c.value) for c in ws["A"] if c.value]
    # No aparece ninguna fila "sin categorizar" para esta cuenta.
    assert not any("413077" in l for l in labels), \
        "413077 sigue itemizada como sin categorizar tras promover+backfill"


# ── AC8: RBAC ───────────────────────────────────────────────────────────────


def test_rbac_family_403_get(tmp_path, monkeypatch):
    _on(monkeypatch)
    client = _make_app(_make_ledger(tmp_path))
    resp = client.get("/api/v1/cuentas-pendientes/", cookies={"access_token": _family()})
    assert resp.status_code == 403


def test_rbac_family_403_promover(tmp_path, monkeypatch):
    _on(monkeypatch)
    client = _make_app(_make_ledger(tmp_path))
    resp = client.post(
        "/api/v1/cuentas-pendientes/413077/promover",
        json={"categoria1": "X", "categoria2": "Y", "categoria3": "Z"},
        cookies={"access_token": _family()},
    )
    assert resp.status_code == 403


def test_rbac_admin_pasa_get(tmp_path, monkeypatch):
    # AC8 nombra admin explícitamente: require_role(["contador","admin"]) lo cubre.
    _on(monkeypatch)
    client = _make_app(_make_ledger(tmp_path))
    resp = client.get("/api/v1/cuentas-pendientes/", cookies={"access_token": _admin()})
    assert resp.status_code == 200
