"""Genera un BORRADOR de balance consolidado (EAG + hijas) desde Beancount
para que el contador lo revise.

Uso:
    PYTHONUTF8=1 python scripts/draft_balance_sheet.py [YYYY-MM-DD]

Fecha de corte por defecto: 2026-05-31 (para comparar con el periodo 05-2026 del
contador). Salida: _handoff/Borrador-Balance-EAG-Consolidado-<fecha>.xlsx

Consolidado = todas las entidades (decisión Ary 2026-06-18: las hijas son egresos
de EAG). El saldo se presenta en su signo natural (Activo deudor +, Pasivo/
Patrimonio acreedor +). Donde el sistema actual (tab Sheets `balance_sheet_eag`)
difiere, se muestra el valor y la diferencia para que el contador adjudique.
"""
from __future__ import annotations

import sys
from collections import defaultdict
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from beancount import loader
from beancount.core.data import Open, Transaction

LEDGER = "ledger/main.beancount"
ROOT_ES = {"Assets": "ACTIVO", "Liabilities": "PASIVO", "Equity": "PATRIMONIO"}

# Saldos del sistema actual (tab Sheets `balance_sheet_eag`) AT 2026-05-31,
# tomados de _handoff/parity-out2.txt — solo las cuentas que diferían.
SHEETS_ACTUAL = {
    "211005": -349482802, "430006": 445889076, "430008": 243641184,
    "430005": 503533518, "430017": 3926309, "430007": 252835403,
    "430009": 44732339, "910021": 1863000000, "710021": 1863000000,
    "610021": 1863000000, "810021": 1863000000, "610022": 784766379,
    "910099": 620158842, "710005": 50513217, "111007": 4320446,
    "113018": -3273174419, "910005": 27199076, "610005": 21962365,
    "810005": 21820162, "111009": 22201191, "115019": 2177595625,
    "115099": 177301862, "613019": 5664366, "111005": 42446187,
    "710006": 1000000, "610006": 1000000, "910006": 1000000,
    "115001": 8142970, "115005": 0, "613005": 60000, "610999": 1685,
}


def build(cutoff: date):
    entries, errs, _ = loader.load_file(LEDGER)
    if errs:
        print(f"⚠️  {len(errs)} errores de carga en el ledger", file=sys.stderr)

    meta = {e.account: (e.meta or {}) for e in entries if isinstance(e, Open)}
    bal: dict[str, float] = defaultdict(float)
    for e in entries:
        if not isinstance(e, Transaction) or e.date > cutoff:
            continue
        for p in e.postings:
            root = p.account.split(":")[0]
            if root in ROOT_ES and p.units and p.units.currency == "CLP":
                bal[p.account] += float(p.units.number)

    rows = []
    for account, raw in bal.items():
        root = account.split(":")[0]
        entity = account.split(":")[1] if account.count(":") >= 1 else ""
        natural = raw if root == "Assets" else -raw  # signo natural
        if abs(natural) < 1:
            continue
        m = meta.get(account, {})
        code = str(m.get("code", ""))
        actual = SHEETS_ACTUAL.get(code)
        rows.append({
            "root": root,
            "raiz_es": ROOT_ES[root],
            "entidad": entity,
            "codigo": code,
            "cuenta": str(m.get("laudus_account_name", account)),
            "cat1": str(m.get("laudus_categoria1", "")),
            "cat2": str(m.get("laudus_categoria2", "")),
            "cat3": str(m.get("laudus_categoria3", "")),
            "saldo": natural,
            "actual": actual,
            "dif": (natural - actual) if actual is not None else None,
        })

    order = {"Assets": 0, "Liabilities": 1, "Equity": 2}
    rows.sort(key=lambda r: (order[r["root"]], r["codigo"]))
    return rows


def write_xlsx(rows, cutoff: date, path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Balance Consolidado"

    bold = Font(bold=True)
    title_f = Font(bold=True, size=14)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    hdr_font = Font(bold=True, color="FFFFFF")
    sub_fill = PatternFill("solid", fgColor="DDEBF7")
    money = "#,##0;(#,##0)"
    right = Alignment(horizontal="right")

    ws["A1"] = "BORRADOR — Balance General Consolidado EAG"
    ws["A1"].font = title_f
    ws["A2"] = f"Entidad: EAG consolidado (EAG + Jocelyn + Jeannette + Johanna + Jael)"
    ws["A3"] = f"Fecha de corte: {cutoff.isoformat()}"
    ws["A4"] = "Fuente: Beancount (derivado de Laudus). Para revisión del contador."
    for r in (2, 3, 4):
        ws[f"A{r}"].font = Font(italic=True, size=10)

    headers = ["Raíz", "Entidad", "Código", "Cuenta", "Categoría 1",
               "Categoría 2", "Saldo (beancount)", "Saldo sistema actual",
               "Diferencia", "Obs."]
    hrow = 6
    for c, h in enumerate(headers, 1):
        cell = ws.cell(hrow, c, h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    r = hrow + 1
    totals = defaultdict(float)
    cur_root = None
    for row in rows:
        if row["root"] != cur_root:
            cur_root = row["root"]
            sec = ws.cell(r, 1, row["raiz_es"])
            sec.font = bold
            for c in range(1, len(headers) + 1):
                ws.cell(r, c).fill = sub_fill
            r += 1
        obs = ""
        if row["dif"] is not None and abs(row["dif"]) >= 1:
            obs = "⚠ difiere del sistema actual"
        if row["saldo"] < 0:
            obs = (obs + "; " if obs else "") + "saldo negativo (revisar)"
        vals = [row["raiz_es"], row["entidad"], row["codigo"], row["cuenta"],
                row["cat1"], row["cat2"], row["saldo"], row["actual"],
                row["dif"], obs]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            if c in (7, 8, 9):
                cell.number_format = money
                cell.alignment = right
        totals[row["root"]] += row["saldo"]
        r += 1

    r += 1
    activos = totals["Assets"]
    pasivos = totals["Liabilities"]
    patrim_cuadre = activos - pasivos
    for label, val in [
        ("TOTAL ACTIVOS", activos),
        ("TOTAL PASIVOS", pasivos),
        ("PATRIMONIO (cuadre = Activos − Pasivos)", patrim_cuadre),
    ]:
        ws.cell(r, 4, label).font = bold
        cell = ws.cell(r, 7, val)
        cell.font = bold
        cell.number_format = money
        cell.alignment = right
        r += 1

    r += 1
    notas = [
        "NOTAS PARA EL CONTADOR:",
        "• Borrador generado desde el ledger Beancount (movimientos importados de Laudus).",
        "• Consolidado = EAG + las 4 hijas (sus inversiones aparecen acá; el detalle por entidad",
        "  está disponible aparte).",
        "• PATRIMONIO se muestra como cuadre (Activos − Pasivos). La apertura ('Saldo anterior')",
        "  hoy está modelada dentro de los movimientos, no como capital separado — definir con el",
        "  contador la estructura de capital / resultados acumulados.",
        "• Columna 'Saldo sistema actual' = tab histórico de Sheets (solo donde difiere). Las",
        "  diferencias son las cuentas a adjudicar; posible staleness del sistema actual.",
        "• Revisar cuentas con saldo negativo (ej. 113018 'Inv. Tecnion Limitada').",
    ]
    for n in notas:
        cell = ws.cell(r, 1, n)
        cell.font = bold if n.endswith(":") else Font(size=10)
        r += 1

    widths = [12, 12, 9, 38, 22, 22, 18, 18, 16, 28]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A7"

    wb.save(path)
    return activos, pasivos, patrim_cuadre


def main():
    arg = sys.argv[1] if len(sys.argv) > 1 else "2026-05-31"
    cutoff = date.fromisoformat(arg)
    rows = build(cutoff)
    out = f"_handoff/Borrador-Balance-EAG-Consolidado-{cutoff.isoformat()}.xlsx"
    a, p, pat = write_xlsx(rows, cutoff, out)
    print(f"✅ {out}")
    print(f"   {len(rows)} cuentas | Activos {a:,.0f} | Pasivos {p:,.0f} | Patrimonio(cuadre) {pat:,.0f}")


if __name__ == "__main__":
    main()
